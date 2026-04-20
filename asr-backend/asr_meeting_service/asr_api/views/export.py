from django.conf import settings
from django.http import FileResponse
from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from docx import Document as DocxDocument
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os
import time
import uuid
from datetime import datetime
import traceback

# ------------------- Word导出接口（转录文本） -------------------
@method_decorator(csrf_exempt, name='dispatch')
class ExportTranscriptionWordView(APIView):
    def post(self, request):
        """导出语音分离结果为Word文档 （转录文本）"""
        try:
            # 1. 提取参数
            transcription_text = request.data.get("transcription_text", "").strip()
            file_name = request.data.get("file_name", "会议记录")
            
            # 2. 输入校验
            if not transcription_text:
                return Response(
                    {"status": "failed", "detail": "语音分离文本不能为空"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # 3. 生成唯一文件名
            file_uuid = str(uuid.uuid4())
            doc_path = os.path.join(settings.TEMP_DIR, f"{file_name}_语音分离_{file_uuid}.docx")
            
            # 4. 创建Word文档
            doc = DocxDocument()
            
            # 标题
            title_para = doc.add_heading(level=0)
            title_run = title_para.add_run(f"{file_name} - 语音分离结果")
            title_run.font.size = Pt(20)
            title_run.bold = True
            title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # 生成时间
            time_para = doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            time_para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
            time_para.add_run("\n\n")
            
            # 内容
            content_para = doc.add_paragraph()
            content_run = content_para.add_run(transcription_text)
            content_run.font.size = Pt(12)
            content_run.font.name = "微软雅黑"
            content_para.line_spacing = 1.5
            
            # 保存文档
            doc.save(doc_path)
            
            # 5. 返回文件下载
            response = FileResponse(
                open(doc_path, 'rb'),
                filename=f"{file_name}_语音分离结果.docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
            # 注：生产环境需添加文件清理逻辑（如定时任务）
            return response
        
        except Exception as e:
            traceback.print_exc()
            return Response(
                {"status": "failed", "detail": f"生成语音分离Word失败：{str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

# ------------------- Word导出接口（会议纪要） -------------------
@method_decorator(csrf_exempt, name='dispatch')
class ExportSummaryWordView(APIView):
    def post(self, request):
        """导出会议纪要为Word文档"""
        try:
            # 1. 提取参数
            summary_text = request.data.get("summary_text", "").strip()
            file_name = request.data.get("file_name", "会议记录")
            
            # 2. 输入校验
            if not summary_text:
                return Response(
                    {"status": "failed", "detail": "会议纪要文本不能为空"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # 3. 生成唯一文件名
            file_uuid = str(uuid.uuid4())
            doc_path = os.path.join(settings.TEMP_DIR, f"{file_name}_会议纪要_{file_uuid}.docx")
            
            # 4. 创建Word文档
            doc = DocxDocument()
            
            # 标题
            title_para = doc.add_heading(level=0)
            title_run = title_para.add_run(f"{file_name} - 会议纪要")
            title_run.font.size = Pt(20)
            title_run.bold = True
            title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # 生成时间
            time_para = doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            time_para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
            time_para.add_run("\n\n")
            
            # 内容
            content_para = doc.add_paragraph()
            content_run = content_para.add_run(summary_text)
            content_run.font.size = Pt(12)
            content_run.font.name = "微软雅黑"
            content_para.line_spacing = 1.5
            
            # 保存文档
            doc.save(doc_path)
            
            # 5. 返回文件下载
            response = FileResponse(
                open(doc_path, 'rb'),
                filename=f"{file_name}_会议纪要.docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
            return response
        
        except Exception as e:
            traceback.print_exc()
            return Response(
                {"status": "failed", "detail": f"生成会议纪要Word失败：{str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

# ------------------- Word导出接口（会议摘要） -------------------
@method_decorator(csrf_exempt, name='dispatch')
class ExportAbstractWordView(APIView):
    def post(self, request):
        """导出会议摘要为Word文档"""
        try:
            # 1. 提取参数
            abstract_text = request.data.get("abstract_text", "").strip()
            file_name = request.data.get("file_name", "会议记录")
            
            # 2. 输入校验
            if not abstract_text:
                return Response(
                    {"status": "failed", "detail": "会议摘要文本不能为空"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # 3. 生成唯一文件名
            file_uuid = str(uuid.uuid4())
            doc_path = os.path.join(settings.TEMP_DIR, f"{file_name}_会议摘要_{file_uuid}.docx")
            
            # 4. 创建Word文档
            doc = DocxDocument()
            
            # 标题
            title_para = doc.add_heading(level=0)
            title_run = title_para.add_run(f"{file_name} - 会议摘要")
            title_run.font.size = Pt(20)
            title_run.bold = True
            title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # 生成时间
            time_para = doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            time_para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
            time_para.add_run("\n\n")
            
            # 摘要内容（特殊格式优化）
            abstract_para = doc.add_paragraph()
            abstract_run = abstract_para.add_run("【核心摘要】\n")
            abstract_run.bold = True
            abstract_run.font.size = Pt(14)
            
            content_run = abstract_para.add_run(abstract_text)
            content_run.font.size = Pt(12)
            content_run.font.name = "微软雅黑"
            abstract_para.line_spacing = 1.5
            
            # 保存文档
            doc.save(doc_path)
            
            # 5. 返回文件下载
            response = FileResponse(
                open(doc_path, 'rb'),
                filename=f"{file_name}_会议摘要.docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
            return response
        
        except Exception as e:
            traceback.print_exc()
            return Response(
                {"status": "failed", "detail": f"生成会议摘要Word失败：{str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )